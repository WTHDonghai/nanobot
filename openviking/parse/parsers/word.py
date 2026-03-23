# Copyright (c) 2026 Beijing Volcano Engine Technology Co., Ltd.
# SPDX-License-Identifier: Apache-2.0
"""
Word document (.docx) parser for OpenViking.

Converts Word documents to Markdown then parses using MarkdownParser.
Inspired by microsoft/markitdown approach.
"""

import re
import zipfile
from pathlib import Path
from typing import List, Optional, Union
from urllib.parse import unquote, urlparse

from openviking.parse.base import ParseResult
from openviking.parse.parsers.base_parser import BaseParser
from openviking_cli.utils.config.parser_config import ParserConfig
from openviking_cli.utils.logger import get_logger

logger = get_logger(__name__)


class WordParser(BaseParser):
    """
    Word document parser for OpenViking.

    Supports: .docx

    Converts Word documents to Markdown using python-docx,
    then delegates to MarkdownParser for tree structure creation.
    """

    EMBEDDED_IMAGES_DIR = "_images"
    IMAGE_PLACEHOLDER_SCHEME = "ov-asset://"
    WORD_CONTROL_CHAR_RE = re.compile(r"[\x01\x13\x14\x15]")
    INCLUDEPICTURE_RE = re.compile(
        r'INCLUDEPICTURE(?:\s+\\d)?\s+"([^"]+)"(?:\s+\\\*\s+MERGEFORMATINET)?',
        re.IGNORECASE,
    )

    def __init__(self, config: Optional[ParserConfig] = None):
        """Initialize Word parser."""
        from openviking.parse.parsers.markdown import MarkdownParser

        self._md_parser = MarkdownParser(config=config)
        self.config = config or ParserConfig()

    @property
    def supported_extensions(self) -> List[str]:
        return [".docx"]

    async def parse(self, source: Union[str, Path], instruction: str = "", **kwargs) -> ParseResult:
        """Parse Word document from file path."""
        path = Path(source)

        if path.exists():
            import docx

            doc = docx.Document(path)
            image_map = self._build_image_reference_map(doc)
            markdown_content = self._convert_to_markdown(doc, image_map)
            result = await self._md_parser.parse_content(
                markdown_content, source_path=str(path), instruction=instruction, **kwargs
            )
            embedded_images = self._extract_embedded_images(path, image_map)
            written_images = await self._write_embedded_images(result.temp_dir_path, embedded_images)
            result.meta = result.meta or {}
            result.meta["embedded_image_count"] = len(written_images)
            if written_images:
                result.meta["embedded_images_dir"] = self.EMBEDDED_IMAGES_DIR
        else:
            result = await self._md_parser.parse_content(
                str(source), instruction=instruction, **kwargs
            )
        result.source_format = "docx"
        result.parser_name = "WordParser"
        return result

    async def parse_content(
        self, content: str, source_path: Optional[str] = None, instruction: str = "", **kwargs
    ) -> ParseResult:
        """Parse content - delegates to MarkdownParser."""
        result = await self._md_parser.parse_content(content, source_path, **kwargs)
        result.source_format = "docx"
        result.parser_name = "WordParser"
        return result

    def _convert_to_markdown(self, doc, image_map: dict[str, str]) -> str:
        """Convert Word document to Markdown string.

        Iterates the document body in order so that tables appear in their
        original position rather than being appended at the end.
        """
        markdown_parts = []

        # Map XML table elements to python-docx Table objects for O(1) lookup
        table_by_element = {table._tbl: table for table in doc.tables}

        # Walk the document body in order to preserve table positions
        from docx.oxml.ns import qn

        for child in doc.element.body:
            if child.tag == qn("w:p"):
                # It's a paragraph
                from docx.text.paragraph import Paragraph

                paragraph = Paragraph(child, doc)
                paragraph_parts = self._build_paragraph_parts(paragraph, image_map)
                paragraph_content = "".join(paragraph_parts).strip()
                if not paragraph_content:
                    continue

                style_name = paragraph.style.name if paragraph.style else "Normal"

                if style_name.startswith("Heading") and paragraph.text.strip():
                    level = self._extract_heading_level(style_name)
                    markdown_parts.append(f"{'#' * level} {paragraph.text}")
                    trailing_images = [
                        part
                        for part in paragraph_parts
                        if part.startswith("![") and self.IMAGE_PLACEHOLDER_SCHEME in part
                    ]
                    markdown_parts.extend(trailing_images)
                else:
                    markdown_parts.append(paragraph_content)

            elif child.tag == qn("w:tbl"):
                # It's a table
                if child in table_by_element:
                    markdown_parts.append(self._convert_table(table_by_element[child]))

        return "\n\n".join(markdown_parts)

    def _extract_heading_level(self, style_name: str) -> int:
        """Extract heading level from style name."""
        try:
            if "Heading" in style_name:
                parts = style_name.split()
                for part in parts:
                    if part.isdigit():
                        return min(int(part), 6)
        except Exception:
            pass
        return 1

    def _convert_paragraph(self, paragraph, image_map: dict[str, str]) -> str:
        """Convert a paragraph into markdown, preserving embedded image positions."""
        return "".join(self._build_paragraph_parts(paragraph, image_map)).strip()

    def _build_paragraph_parts(self, paragraph, image_map: dict[str, str]) -> list[str]:
        """Build paragraph markdown parts in the original run order."""
        parts = []
        for run in paragraph.runs:
            formatted_text = self._format_run_text(run)
            if formatted_text:
                parts.append(formatted_text)
            parts.extend(self._extract_run_image_refs(run, image_map))
        return parts

    @staticmethod
    def _format_run_text(run) -> str:
        """Convert a run's text formatting into markdown."""
        text = WordParser.WORD_CONTROL_CHAR_RE.sub("", run.text or "")
        text = WordParser.INCLUDEPICTURE_RE.sub("", text).strip()
        if not text:
            return ""
        if run.bold:
            text = f"**{text}**"
        if run.italic:
            text = f"*{text}*"
        if run.underline:
            text = f"<ins>{text}</ins>"
        return text

    def _convert_table(self, table) -> str:
        """Convert Word table to markdown format."""
        if not table.rows:
            return ""

        rows = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            rows.append(row_data)

        from openviking.parse.base import format_table_to_markdown

        return format_table_to_markdown(rows, has_header=True)

    def _extract_embedded_images(
        self, path: Path, image_map: dict[str, str]
    ) -> list[tuple[str, bytes]]:
        """Extract embedded DOCX images from the ZIP package."""
        image_entries: list[tuple[str, bytes]] = []
        known_parts = set(image_map.keys())

        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if not name.startswith("word/media/") or name.endswith("/"):
                    continue

                partname = f"/{name}"
                filename = image_map.get(partname)
                if not filename:
                    if known_parts:
                        continue
                    filename = self._sanitize_asset_name(Path(name).name)
                image_entries.append((filename, archive.read(name)))

        return image_entries

    def _build_image_reference_map(self, doc) -> dict[str, str]:
        """Map DOCX relationship targets to stable extracted filenames."""
        image_map: dict[str, str] = {}
        seen_names: set[str] = set()

        for rel_id, part in doc.part.related_parts.items():
            partname = str(getattr(part, "partname", ""))
            content_type = getattr(part, "content_type", "")
            if not partname.startswith("/word/media/") and not content_type.startswith("image/"):
                continue

            image_map[partname] = self._dedupe_asset_name(Path(partname).name, seen_names)

        return image_map

    def _extract_run_image_refs(self, run, image_map: dict[str, str]) -> list[str]:
        """Extract Markdown image placeholders from a run's drawing elements."""
        refs = []
        seen: set[str] = set()
        blips = run.element.xpath('.//*[local-name()="blip"]')
        for blip in blips:
            rel_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            ref = self._render_relationship_image_ref(run, rel_id, image_map)
            if ref and ref not in seen:
                refs.append(ref)
                seen.add(ref)

            link_rel_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link")
            ref = self._render_relationship_image_ref(run, link_rel_id, image_map)
            if ref and ref not in seen:
                refs.append(ref)
                seen.add(ref)

        image_datas = run.element.xpath('.//*[local-name()="imagedata"]')
        for image_data in image_datas:
            rel_id = (
                image_data.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                or image_data.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                or image_data.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link")
                or image_data.get("{urn:schemas-microsoft-com:office:office}relid")
            )
            ref = self._render_relationship_image_ref(run, rel_id, image_map)
            if ref and ref not in seen:
                refs.append(ref)
                seen.add(ref)

        for target in self._extract_includepicture_targets(run.text or ""):
            if not target.startswith(("http://", "https://")):
                continue
            alt_text = self._image_alt_text_from_target(target)
            ref = f"![{alt_text}]({target})"
            if ref not in seen:
                refs.append(ref)
                seen.add(ref)

        return refs

    def _render_relationship_image_ref(self, run, rel_id: Optional[str], image_map: dict[str, str]) -> Optional[str]:
        """Render one image relationship into a markdown image ref."""
        if not rel_id:
            return None

        rel = run.part.rels.get(rel_id)
        if rel is not None and getattr(rel, "is_external", False):
            target_ref = getattr(rel, "target_ref", "")
            if target_ref.startswith(("http://", "https://")):
                alt_text = self._image_alt_text_from_target(target_ref)
                return f"![{alt_text}]({target_ref})"

        image_part = run.part.related_parts.get(rel_id)
        if not image_part:
            return None

        partname = str(getattr(image_part, "partname", ""))
        filename = image_map.get(partname)
        if not filename:
            return None

        alt_text = Path(filename).stem or "image"
        return f"![{alt_text}]({self.IMAGE_PLACEHOLDER_SCHEME}{filename})"

    @classmethod
    def _extract_includepicture_targets(cls, text: str) -> list[str]:
        """Extract INCLUDEPICTURE targets from raw Word field text."""
        cleaned = cls.WORD_CONTROL_CHAR_RE.sub("", text or "")
        return [match.group(1).strip() for match in cls.INCLUDEPICTURE_RE.finditer(cleaned)]

    @staticmethod
    def _image_alt_text_from_target(target: str) -> str:
        """Build a stable alt text from an external image target."""
        parsed = urlparse(target)
        filename = Path(unquote(parsed.path)).name
        return Path(filename).stem or "image"

    async def _write_embedded_images(
        self, temp_uri: Optional[str], image_entries: list[tuple[str, bytes]]
    ) -> list[str]:
        """Persist extracted images under the parsed document root."""
        if not temp_uri or not image_entries:
            return []

        viking_fs = self._get_viking_fs()
        entries = await viking_fs.ls(temp_uri)
        doc_dirs = [e for e in entries if e.get("isDir") and e["name"] not in {".", ".."}]

        if len(doc_dirs) != 1:
            logger.warning(
                f"[WordParser] Expected 1 document directory in {temp_uri}, found {len(doc_dirs)}"
            )
            return []

        doc_root_uri = f"{temp_uri}/{doc_dirs[0]['name']}"
        images_dir_uri = f"{doc_root_uri}/{self.EMBEDDED_IMAGES_DIR}"
        await viking_fs.mkdir(images_dir_uri, exist_ok=True)

        written_images: list[str] = []
        for filename, image_data in image_entries:
            await viking_fs.write(f"{images_dir_uri}/{filename}", image_data)
            written_images.append(filename)

        return written_images

    def _dedupe_asset_name(self, filename: str, seen_names: set[str]) -> str:
        """Create a filesystem-safe unique asset filename."""
        sanitized = self._sanitize_asset_name(filename)
        if sanitized not in seen_names:
            seen_names.add(sanitized)
            return sanitized

        stem = Path(sanitized).stem or "image"
        suffix = Path(sanitized).suffix
        index = 1
        while True:
            candidate = f"{stem}_{index}{suffix}"
            if candidate not in seen_names:
                seen_names.add(candidate)
                return candidate
            index += 1

    def _sanitize_asset_name(self, filename: str) -> str:
        """Normalize embedded asset filenames for VikingFS paths."""
        basename = Path(filename).name
        sanitized = re.sub(r"[^A-Za-z0-9._-]+", "_", basename).strip("._")
        suffix = Path(basename).suffix.lower()
        if not sanitized:
            sanitized = f"image{suffix or '.bin'}"
        elif suffix and not sanitized.endswith(suffix):
            sanitized = f"{sanitized}{suffix}"
        return sanitized
