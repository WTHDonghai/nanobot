from pathlib import Path
import sys

from docx import Document


def iter_blocks(doc: Document):
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            yield ("P", text)

    for idx, table in enumerate(doc.tables, 1):
        yield ("T", f"[Table {idx}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if any(cells):
                yield ("R", " | ".join(cells))


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: extract_docx_text.py <docx> [<docx> ...]", file=sys.stderr)
        return 2

    for path_str in sys.argv[1:]:
        path = Path(path_str)
        print(f"\n===== {path.name} =====")
        doc = Document(path)
        for kind, text in iter_blocks(doc):
            print(f"{kind}: {text}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
