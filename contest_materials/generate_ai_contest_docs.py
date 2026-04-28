from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "contest_materials" / "output"
LOGO = ROOT / "docs" / "images" / "ov-logo.png"

FONT_CN = "Heiti SC"
FONT_EN = "Arial"
ACCENT = "1F4E79"
ACCENT_DARK = "163957"
ACCENT_LIGHT = "EAF3FA"
MUTED = "F5F7FA"
BORDER = "D9E2EC"


def set_run_font(run, size=None, bold=None, color=None, font=FONT_CN):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)


def style_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size, color in [
        ("Title", 24, ACCENT_DARK),
        ("Heading 1", 16, ACCENT_DARK),
        ("Heading 2", 13, ACCENT),
        ("Heading 3", 11.5, ACCENT_DARK),
    ]:
        style = styles[style_name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)


def add_footer(doc: Document, text: str):
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(text)
        set_run_font(run, size=8.5, color="7B8794")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_padding(cell, top=90, left=120, bottom=90, right=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, header=True, widths=None, font_size=9.2):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_padding(cell)
            if header and r_idx == 0:
                shade_cell(cell, ACCENT)
            elif r_idx % 2 == 0:
                shade_cell(cell, "FAFBFC")
            else:
                shade_cell(cell, "FFFFFF")
            if widths and c_idx < len(widths):
                cell.width = Cm(widths[c_idx])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        bold=(header and r_idx == 0),
                        color=("FFFFFF" if header and r_idx == 0 else "1F2933"),
                    )


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row_data in rows:
        row = table.add_row().cells
        for i, value in enumerate(row_data):
            row[i].text = str(value)
    style_table(table, widths=widths, font_size=font_size)
    doc.add_paragraph()
    return table


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.space_after = Pt(5)
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, bold=True, color=ACCENT_DARK)
        rest = text[len(bold_prefix) :]
        if rest:
            run = p.add_run(rest)
            set_run_font(run)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Cm(0.75 + 0.35 * level)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=10)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=10)


def add_callout(doc, title, body, fill=ACCENT_LIGHT):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_border(cell, color="B7D4EA")
    set_cell_padding(cell, top=150, left=180, bottom=150, right=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    set_run_font(run, size=11, bold=True, color=ACCENT_DARK)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(body)
    set_run_font(run, size=10, color="1F2933")
    doc.add_paragraph()


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, "F2F4F7")
    set_cell_border(cell, color="CCD6E0")
    set_cell_padding(cell, top=120, left=140, bottom=120, right=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, size=8.2, font="Maple Mono NF CN", color="1F2933")
    doc.add_paragraph()


def add_cover(doc, title, subtitle, meta_rows):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(LOGO), width=Inches(1.3))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run(title)
    set_run_font(run, size=24, bold=True, color=ACCENT_DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_run_font(run, size=13, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    run = p.add_run("第二届西软 AI 提效大赛参赛材料")
    set_run_font(run, size=11, bold=True, color="4A5568")
    doc.add_paragraph()
    add_table(doc, ["项目", "内容"], meta_rows, widths=[3.3, 11.8], font_size=9.6)
    add_callout(
        doc,
        "一句话定位",
        "以 OpenViking 为通用 AI 知识库与上下文数据库底座，把企业分散的文档、经验、资质、案例和会话记忆组织成可检索、可观察、可复用的 AI 能力层，并通过智能客服与智能投标验证落地价值。",
    )
    doc.add_page_break()


def create_solution_doc():
    doc = Document()
    style_document(doc)
    add_footer(doc, "OpenViking 通用 AI 知识库底座参赛方案")
    add_cover(
        doc,
        "OpenViking 通用 AI 知识库底座",
        "赋能智能客服与智能投标的可复用上下文数据库方案",
        [
            ("推荐赛道", "赛道二：客户 AI 赋能；同时兼具赛道一内部提效价值"),
            ("参赛作品形态", "AI 应用解决方案 + 可运行知识库底座 + 客服机器人案例 + 智能投标案例"),
            ("提交阶段", "报名/初赛材料初稿"),
            ("建议团队", "研发、实施、客服、售前/投标联合组队；成员信息待填写"),
            ("版本日期", "2026 年 5 月"),
        ],
    )

    doc.add_heading("一、参赛定位与材料要求映射", level=1)
    add_para(
        doc,
        "本项目建议以“客户 AI 赋能”赛道参赛：OpenViking 本身不是单一客服机器人，而是可复用的企业 AI 知识库底座。客服机器人是已落地的成功场景，智能投标是同一底座向售前与交付场景扩展的第二个高价值应用。该定位更符合大赛强调的“优秀作品转化为可复用产品或解决方案”的目标。",
    )
    add_table(
        doc,
        ["官方材料要求", "本材料对应内容", "完成状态"],
        [
            ("项目背景", "第二章：业务痛点、影响、AI 解决原因", "已覆盖"),
            ("目标设定", "第三章：可衡量目标与验证方法", "已覆盖"),
            ("解决方案概述", "第四章：总体架构、数据流、工作流", "已覆盖"),
            ("详细设计", "第五章：模型、数据、检索、集成、安全", "已覆盖"),
            ("预期效果评估", "第八章：效率、质量、成本、推广价值", "已覆盖"),
            ("部署/技术文档", "第七章与附录：部署路径、配置、接口、测试", "已覆盖"),
            ("商业价值分析", "第九章概要；另有独立商业价值报告", "已覆盖"),
            ("演示视频", "另有《演示视频脚本》文档，控制在 5 分钟内", "已覆盖"),
        ],
        widths=[4.2, 7.5, 3.2],
    )

    doc.add_heading("二、项目背景", level=1)
    add_para(
        doc,
        "酒店、景区、文旅和企业服务场景中，知识密集型工作正在快速增加：客户咨询依赖产品手册和 FAQ，实施交付依赖历史工单和操作经验，售前投标依赖资质证照、产品方案、案例资料和交付证据。传统知识库通常停留在“文档存储 + 关键词搜索”层面，难以直接成为 AI 应用可用的上下文。",
    )
    add_table(
        doc,
        ["核心问题", "业务影响", "为什么适合用 AI 解决"],
        [
            ("知识碎片化", "资料分散在文档、网盘、聊天记录、工单和个人经验中，复用成本高。", "大模型擅长理解自然语言，但需要稳定、可追溯、结构化的上下文输入。"),
            ("检索命中不稳定", "传统 RAG 扁平切片容易丢失目录语义，回答缺少全局背景。", "OpenViking 用文件系统范式和目录递归检索，让 AI 先定位知识区域，再深入读取证据。"),
            ("材料生产重复", "客服答复、投标章节、方案说明大量重复劳动，且质量依赖个人经验。", "AI 可基于统一知识库生成初稿、引用证据并暴露缺口，把人工精力转向审核和决策。"),
            ("结果难以审计", "AI 回答如果没有来源和检索轨迹，业务团队不敢直接使用。", "OpenViking 保留 URI、检索路径和证据包，便于追溯、纠错和持续优化。"),
        ],
        widths=[3.2, 5.9, 5.9],
    )

    doc.add_heading("三、目标设定", level=1)
    add_para(
        doc,
        "目标分为两类：一类是已有项目资料可以支撑的技术验证目标，另一类是参赛演示与后续试点中建议量化跟踪的业务目标。为避免夸大，本方案将“已验证数据”和“试点目标”分开表述。",
    )
    add_table(
        doc,
        ["目标类别", "量化目标", "验证方式"],
        [
            ("已验证：上下文效果", "在公开评测记录中，OpenViking 结合 OpenClaw 后任务完成率相对原生记忆提升约 43% - 49%。", "基于项目 README_CN 中 LoCoMo10 评测记录。"),
            ("已验证：Token 成本", "同一评测中输入 token 成本相对原生记忆降低约 83% - 91%，相对 LanceDB 方案降低约 92% - 96%。", "基于项目 README_CN 中实验组对比。"),
            ("初赛演示目标", "5 分钟内展示一个知识库底座支撑两个应用：客服问答与投标证据收集。", "演示视频脚本按 0:00 - 5:00 编排。"),
            ("客服试点目标", "知识沉淀完成后，重复咨询人工处理时间目标降低 30% - 50%，常见问题自助命中率目标达到 60% 以上。", "通过客服日志、人工接管率、问题命中率统计。"),
            ("投标试点目标", "资质/方案/案例检索从人工翻找缩短到分钟级，投标材料初稿准备时间目标降低 40% 以上。", "通过章节证据收集耗时、返工次数和人工审核意见统计。"),
        ],
        widths=[3.0, 7.6, 4.4],
    )

    doc.add_heading("四、解决方案概述", level=1)
    add_callout(
        doc,
        "核心思路",
        "把 OpenViking 建设为“企业 AI 大脑的上下文层”：资源、记忆、技能统一进入 VikingFS；系统自动生成 L0/L1/L2 三层上下文；检索时结合目录定位、语义搜索、关键词匹配和证据读取；上层应用通过 Agent 工具或 MCP 接口复用同一知识库。",
    )
    add_code_block(
        doc,
        "企业文档 / FAQ / 工单 / 资质证书 / 案例库 / 会话记录\n"
        "        ↓\n"
        "Parser 文档解析与结构化\n"
        "        ↓\n"
        "VikingFS 分层组织：Resources / Memories / Skills\n"
        "        ↓\n"
        "L0 摘要 + L1 概览 + L2 原文内容\n"
        "        ↓\n"
        "Dense + Sparse 索引，目录递归检索，可视化检索轨迹\n"
        "        ↓\n"
        "Agent 工具 / MCP / HTTP API\n"
        "        ↓\n"
        "西软 AI 客服  |  智能投标助手  |  更多垂直 AI 应用",
    )
    doc.add_page_break()
    doc.add_heading("4.1 架构分层", level=2)
    add_table(
        doc,
        ["层级", "组件", "作用"],
        [
            ("应用层", "西软 AI 客服、智能投标助手、VikingBot", "面向业务用户提供问答、证据收集、材料生成、人工接管等能力。"),
            ("服务层", "OpenViking Server、Session Manager、Tool Registry", "提供 API、会话记忆、工具编排、权限和多租户隔离。"),
            ("知识层", "VikingFS、L0/L1/L2、URI", "将企业知识按文件系统方式组织，保留目录语义和来源。"),
            ("检索层", "Dense + Sparse、目录递归检索、grep/glob/read", "兼顾语义相似、关键词精确、路径定位和证据读取。"),
            ("存储层", "AGFS、VectorDB、RocksDB/Milvus/VikingDB", "提供持久化、索引和可部署扩展能力。"),
        ],
        widths=[2.4, 5.1, 7.5],
    )

    doc.add_heading("五、详细设计", level=1)
    doc.add_heading("5.1 模型与检索设计", level=2)
    add_bullets(
        doc,
        [
            "Embedding 模型：用于文档、图片摘要、FAQ、投标材料的向量化检索，当前配置支持火山引擎 Doubao embedding，也可替换为兼容模型。",
            "VLM/LLM：用于生成 L0 摘要、L1 概览、意图分析、问答生成、投标证据摘要和会话记忆提取。",
            "检索策略：先通过语义和关键词定位高相关目录，再在目录内精细检索并读取原文，避免传统 RAG 只返回孤立切片。",
            "可观察性：每个结果保留 viking:// URI、命中理由、层级、来源文档和必要图片，便于人工核验。",
        ]
    )
    doc.add_heading("5.2 数据来源与预处理", level=2)
    add_table(
        doc,
        ["数据类型", "处理方式", "典型用途"],
        [
            ("产品手册、FAQ、实施文档", "解析为 Markdown/结构化目录，生成摘要和概览，建立索引。", "客服问答、实施支持、售后自助。"),
            ("历史工单、对话记录", "按账号/用户隔离，抽取长期记忆和高频问题。", "相似问题推荐、多轮对话连续性。"),
            ("资质证书、授权资料、案例库", "保留图片证据，支持证书优先排序和本地图片物化。", "投标资格章节、商务证明材料。"),
            ("技术方案、产品参数、架构图", "文本与相关图片共同检索，输出可引用证据片段。", "投标技术方案、售前方案初稿。"),
            ("Skills 与 Prompt 模板", "作为可复用技能挂载到知识库，供不同 Agent 调用。", "客服话术、投标流程、审核标准。"),
        ],
        widths=[3.7, 6.2, 5.1],
    )
    doc.add_heading("5.3 与现有工作模式集成", level=2)
    add_numbered(
        doc,
        [
            "业务人员上传文档或指定资料目录，OpenViking 自动解析、生成摘要并建立索引。",
            "客服、售前、实施等角色通过 Web 页面、Bot、企业 IM 或 MCP 客户端调用同一知识库。",
            "AI 输出带来源的回答、证据包或材料初稿，人工负责确认、修订和发布。",
            "会话结束后，系统将高价值问答、用户偏好和任务经验沉淀为长期记忆，进入下一轮复用。",
        ]
    )
    doc.add_heading("5.4 安全与治理", level=2)
    add_bullets(
        doc,
        [
            "账号、用户和资源空间隔离：支持 account_id、user_id、owner_space 等范围控制。",
            "部署密钥集中配置：模型 API Key、root_api_key、public_bot secret 通过 ov.conf 管理。",
            "人工接管与审核：客服场景保留人工接管，投标场景要求证据先行、人工定稿。",
            "可追溯输出：所有关键材料输出都应保留来源 URI、引用证据和缺口说明。",
        ]
    )

    doc.add_heading("六、核心落地案例", level=1)
    doc.add_heading("6.1 西软 AI 客服", level=2)
    add_para(
        doc,
        "当前项目已经包含“西软 AI 客服”公共测试页与知识召回测试能力。客服机器人可基于产品文档、FAQ 和历史工单回答用户问题，在无法确认时引导人工接管，并将对话沉淀为可复用记忆。",
    )
    add_table(
        doc,
        ["功能", "说明", "业务价值"],
        [
            ("文档问答", "基于产品手册和 FAQ 进行精准回答。", "减少重复咨询，提高一线响应速度。"),
            ("召回测试", "通过 Recall Test 页面检查文档、摘要、命中分数和检索结果。", "让知识库调优有依据，不再靠感觉。"),
            ("多轮记忆", "会话记录可提交到 OpenViking，形成长期上下文。", "支持连续服务，减少用户重复描述。"),
            ("匿名访客", "支持 public_bot 配置，适合客户试用或公开演示。", "降低试点门槛，便于业务推广。"),
        ],
        widths=[3.0, 6.5, 5.5],
    )
    doc.add_heading("6.2 智能投标助手", level=2)
    add_para(
        doc,
        "智能投标场景验证了 OpenViking 不是只服务客服问答的单点应用，而是能为复杂知识工作提供“证据检索 + 缺口识别 + 初稿生成”的通用底座。项目中已经提供投标材料服务、Agent 工具和 MCP Server。",
    )
    add_table(
        doc,
        ["工具/接口", "能力", "输出"],
        [
            ("search_certificates", "搜索营业执照、资质证书、授权资料等证据。", "带图片的证据项和来源 URI。"),
            ("search_solution_materials", "搜索方案、产品参数、案例和架构资料。", "摘录、相关图片和可引用 Markdown。"),
            ("collect_bid_evidence", "按投标章节需求收集证据并识别缺口。", "EvidencePack、claims coverage、gaps。"),
            ("openviking_read/search/list/glob", "开放底层文件系统检索能力给外部 Agent。", "可被 Hermes、Codex 等 MCP 客户端复用。"),
        ],
        widths=[4.0, 6.2, 4.8],
    )

    doc.add_page_break()
    doc.add_heading("七、实现路径与部署说明", level=1)
    add_table(
        doc,
        ["阶段", "工作内容", "交付物"],
        [
            ("第 1 阶段：知识导入", "整理客服 FAQ、产品手册、资质证书、案例库，导入 OpenViking。", "可检索知识库、文档摘要、目录结构。"),
            ("第 2 阶段：客服试点", "接入 public_bot 与客服页面，选择典型问题集做命中率评估。", "客服演示环境、问题集、召回报告。"),
            ("第 3 阶段：投标试点", "接入投标资料，按章节验证证据包和缺口识别。", "投标证据包、章节初稿、缺口清单。"),
            ("第 4 阶段：产品化", "沉淀模板、权限、监控、部署脚本与客户化配置。", "可复制私有化部署方案。"),
        ],
        widths=[3.3, 7.5, 4.2],
    )
    add_table(
        doc,
        ["部署项", "建议配置"],
        [
            ("运行方式", "Docker Compose 双容器：openviking-server 提供 HTTP API、admin、guest；vikingbot 提供 bot gateway。"),
            ("端口", "OpenViking 默认 1933；VikingBot 默认 18790，容器网络内互通。"),
            ("配置文件", "deploy/ov.conf：server、storage、embedding、vlm、bot、log 等配置集中管理。"),
            ("模型", "Embedding 与 VLM 默认可配置为火山引擎兼容模型，也可替换为企业认可的模型服务。"),
            ("健康检查", "curl http://127.0.0.1:1933/health；bot 健康检查 /bot/v1/health。"),
        ],
        widths=[3.2, 11.8],
    )

    doc.add_heading("八、预期效果评估", level=1)
    add_table(
        doc,
        ["维度", "预期/已验证效果", "度量指标"],
        [
            ("效率提升", "客服重复问题自动答复，投标材料检索从人工翻找变为分钟级证据收集。", "平均响应时长、人工接管率、章节准备耗时。"),
            ("质量提升", "回答与投标材料带来源 URI 和证据，降低凭记忆编写导致的错误。", "引用覆盖率、审核返工次数、缺口关闭率。"),
            ("成本降低", "上下文按 L0/L1/L2 分层按需加载，减少无效 token。", "单次任务输入 token、模型调用成本。"),
            ("创新价值", "同一知识库可支撑客服、投标、法务、研发、HR 等多个 Agent。", "新增应用接入数量、复用资料数量。"),
            ("已验证数据", "LoCoMo10 评测中，OpenViking 方案任务完成率相对基线提升约 43% - 49%，输入 token 成本降低约 83% - 91%。", "README_CN 中实验记录。"),
        ],
        widths=[2.8, 8.5, 3.7],
    )

    doc.add_heading("九、商业价值分析概要", level=1)
    add_para(
        doc,
        "对于客户 AI 赋能赛道，OpenViking 的商业价值不止于“做一个机器人”，而是把客户的业务知识资产转化为可运营、可复用、可审计的 AI 能力。独立商业价值分析报告已按官方规范另行整理。",
    )
    add_table(
        doc,
        ["价值方向", "说明"],
        [
            ("客户价值", "7x24 智能服务、知识复用、投标/售前效率提升、经验沉淀。"),
            ("目标客户", "酒店集团、景区、文旅综合体、系统集成客户、内部售前与客服团队。"),
            ("收益模式", "知识库底座、AI 客服模块、投标助手模块、私有化部署、运维服务。"),
            ("竞争优势", "Agent-native 上下文数据库、目录递归检索、三层上下文、MCP 兼容、证据可追溯。"),
            ("推广可行性", "已有 Docker 部署、客服页面、投标 MCP 工具和可验证评测数据。"),
        ],
        widths=[3.2, 11.8],
    )

    doc.add_heading("十、风险与应对", level=1)
    add_table(
        doc,
        ["风险", "影响", "应对策略"],
        [
            ("资料质量不齐", "回答不完整或投标证据不足。", "导入前清洗资料；输出 gaps 缺口清单；建立知识维护机制。"),
            ("模型幻觉", "误答或生成不可靠材料。", "强制来源引用；客服低置信度转人工；投标由人工定稿。"),
            ("权限与隐私", "不同客户/部门资料混用。", "多租户隔离、账号空间隔离、密钥和日志治理。"),
            ("推广阻力", "业务人员担心替代或使用复杂。", "以“辅助和提效”定位，提供简单 Web/Bot 入口和可解释证据。"),
            ("成本波动", "模型调用成本不可控。", "分层上下文、缓存、按需读取、灰度使用。"),
        ],
        widths=[3.0, 5.0, 7.0],
    )

    doc.add_heading("十一、演示与提交计划", level=1)
    add_bullets(
        doc,
        [
            "报名阶段：提交报名表、本方案文档初稿、项目简介。",
            "初赛阶段：补充演示视频链接与提取码，视频控制在 5 分钟内。",
            "决赛阶段：准备 15 分钟现场演示，按“底座能力 + 客服案例 + 投标案例 + 商业价值”讲述。",
            "建议演示数据：选择 3 - 5 份产品资料、2 - 3 个客服高频问题、1 个投标章节需求、若干资质/方案证据。",
        ]
    )
    add_callout(
        doc,
        "最终建议",
        "参赛叙事不要只讲“客服机器人”，而要讲“一个通用知识库底座如何快速孵化多个 AI 应用”。这更贴合创新性、实用性、可推广性三项高权重评分维度。",
        fill="FFF7E6",
    )

    doc.add_page_break()
    doc.add_heading("附录：源码与材料依据", level=1)
    add_table(
        doc,
        ["依据", "位置/说明"],
        [
            ("OpenViking 定位与核心能力", "README_CN.md、DESIGN_DOC.md"),
            ("L0/L1/L2 与目录递归检索", "README_CN.md、docs/zh/concepts/03-context-layers.md、07-retrieval.md"),
            ("客服机器人", "bot/README_CN.md、admin/src/pages/TestBot.tsx、RecallTest.tsx"),
            ("智能投标工具", "bot/vikingbot/services/bid_material.py、agent/tools/bid_material.py、mcp/bid_material_server.py"),
            ("部署方式", "deploy/README.md、deploy/docker-compose.yml、deploy/ov.conf.example"),
            ("大赛要求", "第二届西软AI大赛方案.docx、第二届西软AI大赛提交文档规范说明.docx"),
        ],
        widths=[4.2, 10.8],
    )

    path = OUT / "OpenViking通用知识库_AI大赛参赛方案文档.docx"
    doc.save(path)
    return path


def create_business_doc():
    doc = Document()
    style_document(doc)
    add_footer(doc, "OpenViking 通用 AI 知识库底座商业价值分析报告")
    add_cover(
        doc,
        "商业价值分析报告",
        "OpenViking 通用 AI 知识库底座",
        [
            ("适用赛道", "赛道二：客户 AI 赋能"),
            ("核心客户价值", "降本、增收、提效、体验升级、知识资产化"),
            ("重点案例", "智能客服、智能投标"),
            ("版本日期", "2026 年 5 月"),
        ],
    )
    doc.add_heading("一、客户价值", level=1)
    add_table(
        doc,
        ["客户痛点", "OpenViking 价值", "可衡量指标"],
        [
            ("客户咨询重复且依赖人工经验", "AI 客服基于统一知识库提供 7x24 自助问答，并保留人工接管。", "自助命中率、平均响应时长、人工接管率。"),
            ("售前/投标资料分散", "投标助手自动检索资质、方案、案例和图片证据，生成可审核证据包。", "章节准备耗时、证据覆盖率、返工次数。"),
            ("知识难以持续沉淀", "会话、材料和技能统一进入上下文数据库，越用越完整。", "新增知识数量、高频问题沉淀率、复用次数。"),
            ("AI 应用难以复制", "底座与应用解耦，同一知识库可支撑多个 Agent。", "新增应用接入周期、跨场景复用率。"),
        ],
        widths=[4.0, 7.0, 4.0],
    )

    doc.add_heading("二、目标客户与场景", level=1)
    add_table(
        doc,
        ["客户类型", "典型场景", "切入方式"],
        [
            ("酒店集团/单体酒店", "客人咨询、会员权益、房型/套餐说明、内部员工知识问答。", "从智能客服与员工助手切入。"),
            ("景区/文旅综合体", "导览问答、投诉处理、活动/票务咨询、运营知识库。", "从游客服务与运营问答切入。"),
            ("系统集成/项目型客户", "投标资料管理、方案复用、资质证明、项目案例检索。", "从投标助手和方案助手切入。"),
            ("公司内部业务团队", "客服、实施、售前、研发知识沉淀与复用。", "先内部试点，再产品化输出。"),
        ],
        widths=[3.5, 7.0, 4.5],
    )

    doc.add_heading("三、收益量化", level=1)
    add_para(
        doc,
        "以下收益分为“已有技术验证”和“业务试点目标”。正式推广前建议在一个客服业务线和一个投标业务线各跑 2 - 4 周试点，收集真实基线数据。",
    )
    add_table(
        doc,
        ["收益项", "量化口径", "说明"],
        [
            ("客服提效", "重复问题人工处理时间目标降低 30% - 50%。", "通过问题分类、自动回复命中率和人工接管率统计。"),
            ("投标提效", "资质/方案/案例材料准备时间目标降低 40% 以上。", "通过章节证据收集耗时和审核返工次数统计。"),
            ("模型成本", "已有评测显示输入 token 成本可降低 83% - 91%。", "来源为项目 README_CN 中 OpenClaw 集成评测。"),
            ("任务质量", "已有评测显示任务完成率相对原生记忆提升约 43% - 49%。", "说明 OpenViking 对长上下文任务有实际效果。"),
            ("商业转化", "可形成知识库底座 + AI 客服 + 投标助手 + 私有化部署的组合方案。", "适合纳入客户 AI 赋能产品路线。"),
        ],
        widths=[3.0, 6.5, 5.5],
    )

    doc.add_heading("四、竞争优势", level=1)
    add_table(
        doc,
        ["对比维度", "传统 RAG/单点机器人", "OpenViking 方案"],
        [
            ("知识组织", "扁平向量切片，目录语义弱。", "文件系统范式，Resources/Memories/Skills 统一组织。"),
            ("上下文加载", "容易一次性塞入大量片段。", "L0/L1/L2 分层，按需读取。"),
            ("检索效果", "主要依赖语义相似度。", "目录定位 + 语义搜索 + grep/glob/read。"),
            ("可审计性", "回答来源和路径不清晰。", "URI、证据包、图片和缺口清单可追溯。"),
            ("可扩展性", "一个机器人一个知识库，复用弱。", "同一底座可快速孵化客服、投标、法务、研发等 Agent。"),
        ],
        widths=[3.0, 5.8, 6.2],
    )

    doc.add_heading("五、落地可行性", level=1)
    add_bullets(
        doc,
        [
            "技术基础已具备：项目包含 OpenViking Server、VikingBot、管理后台、客服页面、召回测试页面、投标材料 MCP 工具。",
            "部署路径清晰：Docker Compose 双容器模式，配置集中在 ov.conf，可支持私有化部署。",
            "客户化成本可控：不同客户主要差异在资料导入、知识目录设计、权限配置和业务提示词。",
            "推广策略建议：先做内部客服/售前试点，形成数据闭环，再包装为客户 AI 赋能方案。",
        ]
    )

    doc.add_heading("六、风险应对", level=1)
    add_table(
        doc,
        ["风险", "解决办法"],
        [
            ("知识库资料过旧", "设置资料 owner 和更新时间；回答展示来源；低置信度不直接决策。"),
            ("客户担心数据安全", "支持私有化部署、账号隔离、密钥配置和访问控制。"),
            ("AI 输出不稳定", "投标和客服都保留人工审核/接管；关键输出必须带证据。"),
            ("试点价值难量化", "试点前定义基线指标：处理时长、命中率、返工率、人工接管率。"),
        ],
        widths=[5.0, 10.0],
    )
    path = OUT / "OpenViking通用知识库_商业价值分析报告.docx"
    doc.save(path)
    return path


def create_registration_doc():
    doc = Document()
    style_document(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    add_footer(doc, "OpenViking 通用 AI 知识库底座报名表")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("第二届西软 AI 大赛报名表")
    set_run_font(run, size=18, bold=True, color=ACCENT_DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("项目：OpenViking 通用 AI 知识库底座")
    set_run_font(run, size=12, color=ACCENT)
    doc.add_paragraph()
    add_table(
        doc,
        ["字段", "内容"],
        [
            ("项目名称", "OpenViking 通用 AI 知识库底座：赋能智能客服与智能投标"),
            ("参赛赛道", "赛道二：客户 AI 赋能；兼具赛道一内部提效价值"),
            ("参赛形式", "团队参赛（建议不超过 5 人，成员信息待填写）"),
            ("团队名称", "待填写"),
            ("团队成员", "待填写：姓名 / 部门 / 角色 / 联系方式"),
            ("项目负责人", "待填写"),
            ("联系邮箱", "待填写"),
            ("项目当前状态", "已有 OpenViking Server、管理后台、VikingBot、客服页面、召回测试、智能投标 MCP 工具。"),
        ],
        widths=[4.0, 11.0],
    )
    doc.add_heading("项目简介", level=1)
    add_para(
        doc,
        "OpenViking 是面向 AI 智能体的通用上下文数据库。本项目以 OpenViking 为企业 AI 知识库底座，将产品文档、FAQ、历史工单、资质证书、方案案例和会话记忆统一组织为可检索、可观察、可复用的知识资产。本次参赛以西软 AI 客服作为已落地案例，并扩展智能投标助手，展示一个底座如何快速孵化多个 AI 应用。",
    )
    doc.add_heading("解决的业务问题", level=1)
    add_bullets(
        doc,
        [
            "客服知识分散，重复问题依赖人工经验，响应效率和一致性难保障。",
            "投标资料分散在证书、方案、案例和图片中，人工收集耗时且容易遗漏。",
            "传统 RAG 难以保留目录语义、检索轨迹和证据来源。",
        ]
    )
    doc.add_heading("预期价值", level=1)
    add_bullets(
        doc,
        [
            "客服场景：提升常见问题自助率，降低重复咨询人工处理时间。",
            "投标场景：缩短资质/方案/案例检索时间，减少材料遗漏和返工。",
            "产品化价值：沉淀可复用 AI 知识库底座，支撑客户 AI 赋能产品线。",
        ]
    )
    doc.add_heading("报名提交材料清单", level=1)
    add_table(
        doc,
        ["材料", "状态"],
        [
            ("报名表", "本文档，成员信息待填写"),
            ("项目方案文档初稿", "已生成《OpenViking通用知识库_AI大赛参赛方案文档》"),
            ("商业价值分析报告", "已生成《OpenViking通用知识库_商业价值分析报告》"),
            ("演示视频脚本", "已生成《OpenViking通用知识库_演示视频脚本》"),
        ],
        widths=[7.0, 8.0],
    )
    path = OUT / "OpenViking通用知识库_AI大赛报名表.docx"
    doc.save(path)
    return path


def create_demo_doc():
    doc = Document()
    style_document(doc)
    add_footer(doc, "OpenViking 通用 AI 知识库底座演示视频脚本")
    add_cover(
        doc,
        "5 分钟演示视频脚本",
        "OpenViking 通用 AI 知识库底座",
        [
            ("视频时长", "不超过 5 分钟"),
            ("演示主线", "一个通用知识库底座，支撑智能客服与智能投标两个应用"),
            ("画面要求", "1280×720 至 1920×1080，画面稳定，声音清楚，MP4，建议不超过 300M"),
            ("网盘信息", "百度网盘 URL + 提取码待录制后补充"),
        ],
    )
    doc.add_heading("一、视频结构", level=1)
    add_table(
        doc,
        ["时间", "画面/操作", "解说要点"],
        [
            ("0:00 - 0:30", "标题页 + 项目一句话定位", "我们不是只做一个客服机器人，而是做一个能复用到多个 AI 应用的通用知识库底座。"),
            ("0:30 - 1:10", "展示后台文档/资源列表与召回测试入口", "企业知识被组织为 Resources、Memories、Skills，并生成 L0 摘要、L1 概览、L2 原文。"),
            ("1:10 - 2:15", "进入西软 AI 客服页面，提问一个产品/FAQ 问题", "客服机器人基于知识库回答，减少重复人工咨询，并可保留会话记忆。"),
            ("2:15 - 3:25", "展示智能投标工具或 MCP 调用结果", "投标助手先检索资质、方案、案例证据，再生成可审核证据包和缺口清单。"),
            ("3:25 - 4:15", "展示检索结果、来源 URI、命中理由或文档读取", "OpenViking 的关键差异是可观察、可追溯，而不是黑盒回答。"),
            ("4:15 - 5:00", "回到价值总结页", "一个底座支撑客服、投标和更多 Agent，具备客户 AI 赋能与产品化潜力。"),
        ],
        widths=[2.5, 5.3, 7.2],
        font_size=8.8,
    )

    doc.add_heading("二、逐段口播稿", level=1)
    add_table(
        doc,
        ["段落", "建议口播"],
        [
            (
                "开场",
                "大家好，我们参赛项目是 OpenViking 通用 AI 知识库底座。它的核心价值不是单点机器人，而是把企业分散的文档、经验、证据和会话记忆沉淀成可复用的 AI 上下文层。",
            ),
            (
                "背景",
                "在客服、实施和投标场景里，很多工作不是没有资料，而是资料分散、检索困难、AI 回答缺少来源。传统 RAG 往往把文档切成平面片段，缺少目录语义和检索轨迹。",
            ),
            (
                "方案",
                "OpenViking 使用类似文件系统的范式组织 Resources、Memories 和 Skills。写入时自动生成 L0 摘要、L1 概览、L2 原文，检索时先定位目录，再深入读取证据，兼顾效果、成本和可追溯性。",
            ),
            (
                "客服案例",
                "这是西软 AI 客服场景。用户提问后，机器人从知识库中召回相关资料并回答；遇到不确定问题可交给人工，后续对话也能沉淀为长期记忆。",
            ),
            (
                "投标案例",
                "同一个知识库底座还能支撑智能投标。系统提供资质检索、方案资料检索和章节证据收集工具，输出带来源的 EvidencePack，并指出材料缺口，方便人工审核和完善。",
            ),
            (
                "价值总结",
                "因此，这个项目的可推广性在于：只要接入新的企业资料和场景提示词，就可以孵化新的 AI 应用。客服机器人是第一个成功案例，智能投标是第二个高价值案例，后续还可扩展到法务、研发、HR 和经营分析。",
            ),
        ],
        widths=[2.5, 12.5],
        font_size=9.2,
    )

    doc.add_page_break()
    doc.add_heading("三、录制准备清单", level=1)
    add_bullets(
        doc,
        [
            "准备 3 - 5 份可公开演示的产品文档/FAQ，避免客户敏感信息。",
            "准备 2 - 3 个客服典型问题，至少一个能展示多轮追问。",
            "准备一个投标章节需求，例如“请收集数据库加密方案相关证据”。",
            "准备若干资质证书、方案文档或案例材料，用于展示证据包和图片引用。",
            "录制前检查服务健康：OpenViking /health、VikingBot /bot/v1/health、前端页面可访问。",
            "录制后补充百度网盘 URL 和提取码到报名/提交材料中。",
        ]
    )
    doc.add_heading("四、答辩备选问答", level=1)
    add_table(
        doc,
        ["问题", "建议回答"],
        [
            ("为什么不直接用普通 RAG？", "普通 RAG 更像平面切片库，OpenViking 重点解决目录语义、三层上下文、检索可观察和记忆自迭代。"),
            ("怎么保证回答可信？", "输出保留来源 URI、证据片段、图片和缺口说明；客服支持人工接管，投标必须人工定稿。"),
            ("是否只能做客服？", "不是。客服是成功案例，智能投标展示了同一底座复用到复杂材料生产场景，后续可扩展更多 Agent。"),
            ("落地成本高吗？", "已有 Docker Compose 部署和配置模板，客户化主要集中在资料导入、权限设计和场景提示词。"),
        ],
        widths=[4.2, 10.8],
    )
    path = OUT / "OpenViking通用知识库_AI大赛演示视频脚本.docx"
    doc.save(path)
    return path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [
        create_solution_doc(),
        create_business_doc(),
        create_registration_doc(),
        create_demo_doc(),
    ]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
