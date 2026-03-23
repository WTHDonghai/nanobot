# Copyright (c) 2026 Beijing Volcano Engine Technology Co., Ltd.
# SPDX-License-Identifier: Apache-2.0

"""Tests for WordParser embedded image extraction."""

import base64
from pathlib import Path
from typing import Any
from unittest.mock import AsyncMock

import pytest
from docx import Document

from openviking.parse.base import NodeType, ResourceNode, create_parse_result
from openviking.parse.parsers.word import WordParser

PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO7Z0WQAAAAASUVORK5CYII="
)


class FakeVikingFS:
    """Minimal VikingFS mock for parser output assertions."""

    def __init__(self):
        self.dirs: list[str] = []
        self.files: dict[str, bytes] = {}
        self._temp_counter = 0

    async def mkdir(self, uri: str, exist_ok: bool = False, **kwargs: Any) -> None:
        if uri not in self.dirs:
            self.dirs.append(uri)

    async def write(self, uri: str, data: Any) -> str:
        if isinstance(data, str):
            data = data.encode("utf-8")
        self.files[uri] = data
        return uri

    async def ls(self, uri: str) -> list[dict[str, Any]]:
        prefix = uri.rstrip("/") + "/"
        children: dict[str, bool] = {}

        for key in list(self.files.keys()) + self.dirs:
            if not key.startswith(prefix):
                continue
            rest = key[len(prefix) :]
            if not rest:
                continue
            name = rest.split("/")[0]
            child_uri = f"{prefix}{name}"
            is_dir = "/" in rest[len(name) :] or child_uri in self.dirs
            children[name] = children.get(name, False) or is_dir

        return [
            {"name": name, "uri": f"{uri.rstrip('/')}/{name}", "isDir": is_dir}
            for name, is_dir in sorted(children.items())
        ]

    def create_temp_uri(self) -> str:
        self._temp_counter += 1
        return f"viking://temp/word_{self._temp_counter}"


class FakeImageElement:
    def __init__(self, attrs: dict[str, str]):
        self.attrs = attrs

    def get(self, key: str):
        return self.attrs.get(key)


class FakeElement:
    def __init__(self, blips=None, image_datas=None):
        self._blips = blips or []
        self._image_datas = image_datas or []

    def xpath(self, expr: str):
        if "blip" in expr:
            return self._blips
        if "imagedata" in expr:
            return self._image_datas
        return []


class FakeRel:
    def __init__(self, is_external: bool = False, target_ref: str = ""):
        self.is_external = is_external
        self.target_ref = target_ref


class FakePart:
    def __init__(self, related_parts=None, rels=None):
        self.related_parts = related_parts or {}
        self.rels = rels or {}


class FakeRun:
    def __init__(self, text: str = "", element=None, part=None):
        self.text = text
        self.bold = False
        self.italic = False
        self.underline = False
        self.element = element or FakeElement()
        self.part = part or FakePart()


@pytest.mark.asyncio
async def test_word_parser_extracts_embedded_images(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "tiny.png"
    image_path.write_bytes(PNG_BYTES)

    docx_path = tmp_path / "manual.docx"
    document = Document()
    document.add_heading("Manual", level=1)
    document.add_paragraph("Device overview")
    document.add_picture(str(image_path))
    document.save(docx_path)

    fake_fs = FakeVikingFS()
    temp_uri = fake_fs.create_temp_uri()
    doc_dir = f"{temp_uri}/manual"
    await fake_fs.mkdir(temp_uri)
    await fake_fs.mkdir(doc_dir)

    parse_result = create_parse_result(
        root=ResourceNode(type=NodeType.ROOT, title="Manual", level=0),
        source_path=str(docx_path),
        source_format="markdown",
        parser_name="MarkdownParser",
        parse_time=0.1,
    )
    parse_result.temp_dir_path = temp_uri

    parser = WordParser()
    parser._md_parser.parse_content = AsyncMock(return_value=parse_result)
    monkeypatch.setattr("openviking.storage.viking_fs.get_viking_fs", lambda: fake_fs)

    result = await parser.parse(docx_path)

    image_uris = [uri for uri in fake_fs.files if uri.startswith(f"{doc_dir}/_images/")]
    assert len(image_uris) == 1
    assert image_uris[0].endswith(".png")
    assert fake_fs.files[image_uris[0]]
    assert result.meta["embedded_image_count"] == 1
    assert result.meta["embedded_images_dir"] == "_images"
    markdown_content = parser._md_parser.parse_content.await_args.args[0]
    assert "ov-asset://" in markdown_content


@pytest.mark.asyncio
async def test_word_parser_preserves_inline_image_order(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "tiny.png"
    image_path.write_bytes(PNG_BYTES)

    docx_path = tmp_path / "inline.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Click ")
    paragraph.add_run().add_picture(str(image_path))
    paragraph.add_run(" to continue")
    document.save(docx_path)

    fake_fs = FakeVikingFS()
    temp_uri = fake_fs.create_temp_uri()
    doc_dir = f"{temp_uri}/inline"
    await fake_fs.mkdir(temp_uri)
    await fake_fs.mkdir(doc_dir)

    parse_result = create_parse_result(
        root=ResourceNode(type=NodeType.ROOT, title="Inline", level=0),
        source_path=str(docx_path),
        source_format="markdown",
        parser_name="MarkdownParser",
        parse_time=0.1,
    )
    parse_result.temp_dir_path = temp_uri

    parser = WordParser()
    parser._md_parser.parse_content = AsyncMock(return_value=parse_result)
    monkeypatch.setattr("openviking.storage.viking_fs.get_viking_fs", lambda: fake_fs)

    await parser.parse(docx_path)

    markdown_content = parser._md_parser.parse_content.await_args.args[0]
    assert markdown_content.index("Click ") < markdown_content.index("ov-asset://")
    assert markdown_content.index("ov-asset://") < markdown_content.index(" to continue")


def test_word_parser_strips_includepicture_field_text_but_keeps_following_note() -> None:
    run = FakeRun(
        text=(
            '\x13 INCLUDEPICTURE "C:\\\\Users\\\\ADMINI~1\\\\AppData\\\\Local\\\\Temp\\\\ksohtml\\\\wps7F76.tmp.jpg" '
            '\\* MERGEFORMATINET \x14\x01\x15注意：扫码登录功能需后台配置。'
        )
    )

    assert WordParser._format_run_text(run) == "注意：扫码登录功能需后台配置。"


def test_word_parser_extracts_http_includepicture_as_markdown_image() -> None:
    parser = WordParser()
    run = FakeRun(
        text=(
            '\x13 INCLUDEPICTURE "https://example.com/assets/scan-login.png?token=abc" '
            '\\* MERGEFORMATINET \x14\x01\x15'
        )
    )

    refs = parser._extract_run_image_refs(run, {})

    assert refs == ["![scan-login](https://example.com/assets/scan-login.png?token=abc)"]


def test_word_parser_extracts_vml_imagedata_relationship() -> None:
    parser = WordParser()
    run = FakeRun(
        element=FakeElement(
            image_datas=[
                FakeImageElement(
                    {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": "rId9"}
                )
            ]
        ),
        part=FakePart(
            related_parts={
                "rId9": type("ImagePart", (), {"partname": "/word/media/image1.png"})(),
            },
            rels={"rId9": FakeRel()},
        ),
    )

    refs = parser._extract_run_image_refs(run, {"/word/media/image1.png": "image1.png"})

    assert refs == ["![image1](ov-asset://image1.png)"]
